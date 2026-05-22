"""Data profile refresh — scan code/data/raw/ and rebuild .metadata_cache/."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from maintenance import (
    FAIL, PASS, WARN,
    RAW_DATA_DIR, METADATA_CACHE,
    CheckResult, make_result,
)


def _build_profile(data_file: Path) -> dict[str, Any]:
    profile: dict[str, Any] = {"mtime": data_file.stat().st_mtime}
    try:
        if data_file.suffix.lower() == ".csv":
            with open(data_file, newline="", encoding="utf-8") as fh:
                reader = csv.reader(fh)
                header = next(reader, None)
                if header is None:
                    profile["row_count"] = 0
                    profile["column_count"] = 0
                    profile["missing_values"] = {}
                    return profile
                rows = list(reader)
                profile["row_count"] = len(rows)
                profile["column_count"] = len(header)
                missing: dict[str, int] = {col: 0 for col in header}
                for row in rows:
                    for i, val in enumerate(row):
                        if i < len(header) and val.strip() == "":
                            missing[header[i]] += 1
                profile["missing_values"] = missing
        elif data_file.suffix.lower() in (".xlsx", ".xlsm"):
            try:
                from openpyxl import load_workbook
            except ImportError:
                profile["row_count"] = 0
                profile["column_count"] = 0
                profile["missing_values"] = {}
                profile["note"] = "openpyxl not installed; skipping detail"
                return profile
            wb = load_workbook(data_file, read_only=True, data_only=True)
            ws = wb.active
            header = [cell for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])]
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            profile["row_count"] = len(rows)
            profile["column_count"] = len(header)
            missing = {str(col): 0 for col in header if col is not None}
            for row in rows:
                for i, val in enumerate(row):
                    if i < len(header) and (val is None or val == ""):
                        missing[str(header[i])] = missing.get(str(header[i]), 0) + 1
            profile["missing_values"] = missing
            wb.close()
        else:
            profile["row_count"] = 0
            profile["column_count"] = 0
            profile["missing_values"] = {}
            profile["note"] = f"unsupported format: {data_file.suffix}"
    except Exception as exc:
        profile["error"] = str(exc)
    profile.setdefault("sample_data", [])
    return profile


def run() -> list[CheckResult]:
    """Re-scan every file under code/data/raw and rewrite .metadata_cache/."""
    results: list[CheckResult] = []
    if not RAW_DATA_DIR.exists():
        results.append(make_result(FAIL, f"raw data dir missing: {RAW_DATA_DIR}"))
        return results

    METADATA_CACHE.mkdir(parents=True, exist_ok=True)
    data_files = sorted(
        [p for p in RAW_DATA_DIR.iterdir()
         if p.is_file() and p.suffix.lower() in (".csv", ".xlsx", ".xlsm")]
    )

    if not data_files:
        results.append(make_result(WARN, "no data files found in raw/"))
        return results

    for f in data_files:
        cache_key = f.stem.lower() + ".json"
        cache_file = METADATA_CACHE / cache_key
        try:
            profile = _build_profile(f)
            with open(cache_file, "w", encoding="utf-8") as fh:
                json.dump(profile, fh, indent=2)
                fh.write("\n")
            results.append(make_result(
                PASS,
                f"{f.name} -> {profile.get('row_count', '?')} rows, "
                f"{profile.get('column_count', '?')} cols",
            ))
        except Exception as exc:
            results.append(make_result(FAIL, f"{f.name}: {exc}"))

    stale = [
        p for p in METADATA_CACHE.iterdir()
        if p.is_file() and p.suffix == ".json"
        and p.stem not in {f.stem.lower() for f in data_files}
    ]
    for p in stale:
        p.unlink()
        results.append(make_result(WARN, f"removed stale cache: {p.name}"))

    return results
